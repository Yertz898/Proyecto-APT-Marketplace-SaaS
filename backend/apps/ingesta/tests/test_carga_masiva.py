"""
Carga masiva del catálogo.

Es el módulo que más va a fallar, porque el archivo lo escribe una persona a
mano. Así que acá el camino feliz es una prueba de muchas: casi todo lo demás
es cómo se comporta con un archivo que no vino como esperábamos.

Las tres reglas que se comprueban una y otra vez:

- Validar **no escribe nada** en el catálogo.
- Una fila mala **no bota a las demás**.
- Lo que no viene en el archivo **no se borra**.
"""

import csv
import io
import zipfile

import pytest
from openpyxl import Workbook
from PIL import Image

from apps.catalogo.models import Producto
from apps.ingesta.models import Carga, EstadoCarga
from apps.precios.models import TramoPrecio

pytestmark = pytest.mark.django_db

LIMITES = "/api/panel/cargas/limites"
HISTORIAL = "/api/panel/cargas"


def como(client, usuario):
    from rest_framework_simplejwt.tokens import AccessToken

    client.defaults["HTTP_AUTHORIZATION"] = f"Bearer {AccessToken.for_user(usuario)}"
    return client


# ── Archivos de prueba ──────────────────────────────────────────────────────

ENCABEZADOS = [
    "codigo",
    "nombre",
    "precio_neto",
    "forma_venta",
    "unidad_medida",
    "controla_stock",
    "stock",
]

FILA = ["NUEVO-1", "Producto nuevo", "4990", "unidad", "unidad", "si", "10"]


def xlsx(encabezados=None, filas=None, nombre="catalogo.xlsx"):
    from django.core.files.uploadedfile import SimpleUploadedFile

    libro = Workbook()
    hoja = libro.active
    hoja.append(encabezados if encabezados is not None else ENCABEZADOS)
    for fila in filas if filas is not None else [FILA]:
        hoja.append(fila)

    memoria = io.BytesIO()
    libro.save(memoria)
    return SimpleUploadedFile(
        nombre,
        memoria.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )


def csv_archivo(
    encabezados=None,
    filas=None,
    separador=";",
    codificacion="utf-8",
    nombre="catalogo.csv",
):
    from django.core.files.uploadedfile import SimpleUploadedFile

    texto = io.StringIO()
    escritor = csv.writer(texto, delimiter=separador, lineterminator="\n")
    escritor.writerow(encabezados if encabezados is not None else ENCABEZADOS)
    for fila in filas if filas is not None else [FILA]:
        escritor.writerow(fila)

    return SimpleUploadedFile(
        nombre, texto.getvalue().encode(codificacion), content_type="text/csv"
    )


def imagen_bytes(tamano=(8, 8)):
    memoria = io.BytesIO()
    Image.new("RGB", tamano, "#7C3AED").save(memoria, format="PNG")
    return memoria.getvalue()


def zip_archivo(entradas, nombre="fotos.zip"):
    """Un .zip con las entradas dadas: {nombre interno: bytes}."""
    from django.core.files.uploadedfile import SimpleUploadedFile

    memoria = io.BytesIO()
    with zipfile.ZipFile(memoria, "w") as comprimido:
        for interno, contenido in entradas.items():
            comprimido.writestr(interno, contenido)

    return SimpleUploadedFile(
        nombre, memoria.getvalue(), content_type="application/zip"
    )


def validar(client, tipo, archivo):
    return client.post(f"/api/panel/cargas/{tipo}/validar", {"archivo": archivo})


def confirmar(client, carga_id, desactivar=False):
    return client.post(
        f"/api/panel/cargas/{carga_id}/confirmar",
        {"desactivarAusentes": desactivar},
        content_type="application/json",
    )


# ── Validar no escribe ──────────────────────────────────────────────────────


def test_validar_no_toca_el_catalogo(client, dueno_a, tienda_a):
    """
    La regla que da sentido a los dos pasos.

    Si validar escribiera, el resumen que ve el dueño sería un informe de algo
    que ya pasó.
    """
    respuesta = validar(como(client, dueno_a), "productos", xlsx())

    assert respuesta.status_code == 201
    assert respuesta.json()["crear"] == 1
    assert not Producto.objects.de_tienda(tienda_a).exists()


def test_validar_deja_la_carga_pendiente(client, dueno_a, tienda_a):
    validar(como(client, dueno_a), "productos", xlsx())

    carga = Carga.objects.de_tienda(tienda_a).get()
    assert carga.estado == EstadoCarga.VALIDADA
    assert carga.nombre_original == "catalogo.xlsx"


def test_confirmar_escribe(client, dueno_a, tienda_a):
    carga_id = validar(como(client, dueno_a), "productos", xlsx()).json()["cargaId"]

    respuesta = confirmar(como(client, dueno_a), carga_id)

    assert respuesta.status_code == 200
    producto = Producto.objects.de_tienda(tienda_a).get()
    assert producto.codigo == "NUEVO-1"
    assert producto.precio_neto == 4990
    assert producto.stock == 10


# ── Cómo viene el archivo ───────────────────────────────────────────────────


def test_las_columnas_pueden_ir_en_otro_orden(client, dueno_a, tienda_a):
    """Nunca se confía en el orden: se mapea por nombre de encabezado."""
    archivo = xlsx(
        encabezados=["nombre", "precio_neto", "codigo"],
        filas=[["Al revés", "7990", "REVES-1"]],
    )
    carga_id = validar(como(client, dueno_a), "productos", archivo).json()["cargaId"]
    confirmar(como(client, dueno_a), carga_id)

    assert Producto.objects.de_tienda(tienda_a).get().codigo == "REVES-1"


def test_los_encabezados_se_normalizan(client, dueno_a, tienda_a):
    """«Código», «codigo» y «CODIGO» son la misma columna."""
    archivo = xlsx(
        encabezados=["Código", "NOMBRE", "Precio Neto"],
        filas=[["ACENTO-1", "Con acentos", "4990"]],
    )
    carga_id = validar(como(client, dueno_a), "productos", archivo).json()["cargaId"]
    confirmar(como(client, dueno_a), carga_id)

    assert Producto.objects.de_tienda(tienda_a).get().codigo == "ACENTO-1"


def test_una_columna_de_mas_no_rompe(client, dueno_a):
    archivo = xlsx(
        encabezados=["codigo", "nombre", "precio_neto", "inventada"],
        filas=[["EXTRA-1", "Con columna de más", "4990", "lo que sea"]],
    )

    assert validar(como(client, dueno_a), "productos", archivo).json()["crear"] == 1


def test_faltar_una_columna_entera_rechaza_la_planilla(
    client, dueno_a, detalles_de_error
):
    """
    Si falta la columna, no se rechaza fila por fila.

    Cada una fallaría por lo mismo, y trescientos errores idénticos no ayudan
    a nadie: lo que hay que decir es qué columna falta.
    """
    archivo = xlsx(
        encabezados=["codigo", "nombre"], filas=[["SIN-PRECIO", "Sin precio"]]
    )

    respuesta = validar(como(client, dueno_a), "productos", archivo)

    assert respuesta.status_code == 400
    assert "precio_neto" in detalles_de_error(respuesta)["archivo"][0]


def test_una_fila_con_el_precio_vacio_se_rechaza(client, dueno_a):
    """La columna está; la celda no. Eso sí es cosa de la fila."""
    archivo = xlsx(
        filas=[["SIN-PRECIO", "Sin precio", "", "unidad", "unidad", "si", "5"]]
    )

    cuerpo = validar(como(client, dueno_a), "productos", archivo).json()

    assert cuerpo["rechazar"] == 1
    assert "precio" in cuerpo["rechazos"][0]["motivo"].lower()


def test_un_archivo_que_no_es_planilla_no_pasa_como_vacio(
    client, dueno_a, detalles_de_error
):
    """
    Lo peor que podía hacer: aceptar basura con un resumen en cero.

    «0 nuevos, 0 rechazados» se parece demasiado a que funcionó.
    """
    archivo = csv_archivo(
        encabezados=["cualquier", "cosa"], filas=[["a", "b"]], nombre="otra.csv"
    )

    respuesta = validar(como(client, dueno_a), "productos", archivo)

    assert respuesta.status_code == 400
    assert "columnas" in detalles_de_error(respuesta)["archivo"][0]


def test_una_fila_mala_no_bota_a_las_demas(client, dueno_a):
    """
    La regla que hace usable el módulo.

    Si una planilla de 300 filas se cayera por la 150, nadie la usaría.
    """
    archivo = xlsx(
        filas=[
            ["BUENA-1", "Buena", "4990", "unidad", "unidad", "si", "5"],
            ["MALA-1", "Sin precio", "", "unidad", "unidad", "si", "5"],
            ["BUENA-2", "Buena también", "7990", "unidad", "unidad", "si", "5"],
        ]
    )

    cuerpo = validar(como(client, dueno_a), "productos", archivo).json()

    assert cuerpo["crear"] == 2
    assert cuerpo["rechazar"] == 1


def test_el_rechazo_dice_en_que_fila(client, dueno_a):
    """Con el número que la persona ve en Excel, no el índice del programa."""
    archivo = xlsx(
        filas=[
            ["BUENA-1", "Buena", "4990", "unidad", "unidad", "si", "5"],
            ["MALA-1", "Sin precio", "", "unidad", "unidad", "si", "5"],
        ]
    )

    rechazo = validar(como(client, dueno_a), "productos", archivo).json()["rechazos"][0]

    assert rechazo["fila"] == 3


def test_las_filas_en_blanco_se_ignoran(client, dueno_a):
    """Excel deja filas vacías al final más seguido de lo que parece."""
    archivo = xlsx(filas=[FILA, ["", "", "", "", "", "", ""]])

    cuerpo = validar(como(client, dueno_a), "productos", archivo).json()

    assert cuerpo["crear"] == 1
    assert cuerpo["rechazar"] == 0


def test_un_codigo_repetido_dentro_de_la_planilla_se_rechaza(client, dueno_a):
    archivo = xlsx(
        filas=[
            ["IGUAL", "Primera", "4990", "unidad", "unidad", "si", "5"],
            ["IGUAL", "Segunda", "7990", "unidad", "unidad", "si", "5"],
        ]
    )

    cuerpo = validar(como(client, dueno_a), "productos", archivo).json()

    assert cuerpo["crear"] == 1
    assert cuerpo["rechazar"] == 1


# ── CSV, que es donde más varía ─────────────────────────────────────────────


def test_el_csv_con_punto_y_coma(client, dueno_a, tienda_a):
    """Es lo que exporta Excel en español."""
    carga_id = validar(
        como(client, dueno_a), "productos", csv_archivo(separador=";")
    ).json()["cargaId"]
    confirmar(como(client, dueno_a), carga_id)

    assert Producto.objects.de_tienda(tienda_a).get().codigo == "NUEVO-1"


def test_el_csv_con_coma(client, dueno_a):
    archivo = csv_archivo(
        encabezados=["codigo", "nombre", "precio_neto"],
        filas=[["COMA-1", "Con comas", "4990"]],
        separador=",",
    )

    assert validar(como(client, dueno_a), "productos", archivo).json()["crear"] == 1


def test_el_csv_en_windows_1252(client, dueno_a, tienda_a):
    """
    Las tildes y las eñes se rompen acá o en ninguna parte.

    Excel en español guarda con esta codificación si no se elige otra.
    """
    archivo = csv_archivo(
        encabezados=["codigo", "nombre", "precio_neto"],
        filas=[["ENE-1", "Cadena de niño con ñ y tildé", "4990"]],
        codificacion="cp1252",
    )
    carga_id = validar(como(client, dueno_a), "productos", archivo).json()["cargaId"]
    confirmar(como(client, dueno_a), carga_id)

    assert "niño" in Producto.objects.de_tienda(tienda_a).get().nombre


def test_la_coma_decimal_se_entiende(client, dueno_a, tienda_a):
    """En Chile se escribe 1,5 y Excel lo exporta así."""
    archivo = csv_archivo(
        encabezados=[
            "codigo",
            "nombre",
            "precio_neto",
            "forma_venta",
            "unidad_medida",
            "stock",
        ],
        filas=[["PESO-1", "Por gramo", "650", "peso", "g", "125,5"]],
    )
    carga_id = validar(como(client, dueno_a), "productos", archivo).json()["cargaId"]
    confirmar(como(client, dueno_a), carga_id)

    assert str(Producto.objects.de_tienda(tienda_a).get().stock) == "125.500"


# ── Qué archivos se rechazan ────────────────────────────────────────────────


def test_un_archivo_con_la_extension_cambiada_se_rechaza(
    client, dueno_a, detalles_de_error
):
    """
    Se valida por contenido, no por el nombre.

    La extensión la escribe quien sube el archivo y no dice nada de lo que hay
    adentro.
    """
    from django.core.files.uploadedfile import SimpleUploadedFile

    disfrazado = SimpleUploadedFile(
        "catalogo.xlsx", imagen_bytes(), content_type="application/vnd.ms-excel"
    )
    respuesta = validar(como(client, dueno_a), "productos", disfrazado)

    assert respuesta.status_code == 400
    assert "archivo" in detalles_de_error(respuesta)


def test_un_xls_antiguo_se_rechaza_explicando(client, dueno_a, detalles_de_error):
    """
    Con los pasos para arreglarlo, no con un «no se acepta» a secas.

    Es el error más común de todos: Excel ofrece ese formato primero.
    """
    from django.core.files.uploadedfile import SimpleUploadedFile

    # La firma de un .xls antiguo (documento compuesto de Microsoft).
    viejo = SimpleUploadedFile(
        "catalogo.xls",
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600,
        content_type="application/vnd.ms-excel",
    )
    respuesta = validar(como(client, dueno_a), "productos", viejo)

    assert respuesta.status_code == 400
    assert "Guardar como" in detalles_de_error(respuesta)["archivo"][0]


def test_un_archivo_demasiado_grande_se_rechaza(
    client, dueno_a, settings, detalles_de_error
):
    settings.TAMANO_MAXIMO_CARGA = 100

    respuesta = validar(como(client, dueno_a), "productos", xlsx())

    assert respuesta.status_code == 400
    assert "archivo" in detalles_de_error(respuesta)


def test_sin_archivo_no_hay_carga(client, dueno_a, detalles_de_error):
    respuesta = como(client, dueno_a).post("/api/panel/cargas/productos/validar", {})

    assert respuesta.status_code == 400
    assert "archivo" in detalles_de_error(respuesta)


# ── Idempotencia y ausentes ─────────────────────────────────────────────────


def test_subir_dos_veces_la_misma_planilla_no_duplica(client, dueno_a, tienda_a):
    """La clave natural es (tienda, código)."""
    for _ in range(2):
        carga_id = validar(como(client, dueno_a), "productos", xlsx()).json()["cargaId"]
        confirmar(como(client, dueno_a), carga_id)

    assert Producto.objects.de_tienda(tienda_a).count() == 1


def test_la_segunda_carga_actualiza(client, dueno_a, tienda_a):
    carga_id = validar(como(client, dueno_a), "productos", xlsx()).json()["cargaId"]
    confirmar(como(client, dueno_a), carga_id)

    otra = xlsx(
        filas=[["NUEVO-1", "Nombre cambiado", "9990", "unidad", "unidad", "si", "3"]]
    )
    cuerpo = validar(como(client, dueno_a), "productos", otra).json()
    confirmar(como(client, dueno_a), cuerpo["cargaId"])

    assert cuerpo["actualizar"] == 1
    producto = Producto.objects.de_tienda(tienda_a).get()
    assert producto.nombre == "Nombre cambiado"
    assert producto.precio_neto == 9990


def test_el_resumen_avisa_qué_quedó_fuera_de_la_planilla(
    client, dueno_a, tienda_a, producto_a
):
    cuerpo = validar(como(client, dueno_a), "productos", xlsx()).json()

    assert producto_a.codigo in cuerpo["ausentes"]


def test_lo_ausente_no_se_borra(client, dueno_a, tienda_a, producto_a):
    """
    Una planilla incompleta no puede vaciar una tienda.

    Es la protección más importante del módulo: alguien sube las novedades del
    mes y no espera perder el resto del catálogo.
    """
    carga_id = validar(como(client, dueno_a), "productos", xlsx()).json()["cargaId"]
    confirmar(como(client, dueno_a), carga_id)

    producto_a.refresh_from_db()
    assert producto_a.activo is True


def test_lo_ausente_se_desactiva_solo_si_se_pide(client, dueno_a, producto_a):
    carga_id = validar(como(client, dueno_a), "productos", xlsx()).json()["cargaId"]
    confirmar(como(client, dueno_a), carga_id, desactivar=True)

    producto_a.refresh_from_db()
    assert producto_a.activo is False
    assert Producto.objects.filter(pk=producto_a.pk).exists()


# ── Confirmar y descartar ───────────────────────────────────────────────────


def test_una_carga_no_se_confirma_dos_veces(client, dueno_a, tienda_a):
    carga_id = validar(como(client, dueno_a), "productos", xlsx()).json()["cargaId"]
    confirmar(como(client, dueno_a), carga_id)

    assert confirmar(como(client, dueno_a), carga_id).status_code == 400
    assert Producto.objects.de_tienda(tienda_a).count() == 1


def test_descartar_deja_constancia_en_vez_de_borrar(client, dueno_a, tienda_a):
    """El historial tiene que mostrar también lo que se decidió no aplicar."""
    carga_id = validar(como(client, dueno_a), "productos", xlsx()).json()["cargaId"]

    respuesta = como(client, dueno_a).delete(f"/api/panel/cargas/{carga_id}")

    assert respuesta.status_code == 204
    carga = Carga.objects.de_tienda(tienda_a).get()
    assert carga.estado == EstadoCarga.DESCARTADA
    assert not Producto.objects.de_tienda(tienda_a).exists()


def test_una_carga_descartada_no_se_confirma(client, dueno_a):
    carga_id = validar(como(client, dueno_a), "productos", xlsx()).json()["cargaId"]
    como(client, dueno_a).delete(f"/api/panel/cargas/{carga_id}")

    assert confirmar(como(client, dueno_a), carga_id).status_code == 400


# ── Tramos ──────────────────────────────────────────────────────────────────


def test_la_carga_de_tramos(client, dueno_a, tienda_a, producto_a):
    archivo = xlsx(
        encabezados=["codigo_producto", "umbral", "unidad_umbral", "precio_neto"],
        filas=[[producto_a.codigo, "12", "cantidad", "4490"]],
        nombre="tramos.xlsx",
    )
    carga_id = validar(como(client, dueno_a), "tramos", archivo).json()["cargaId"]
    confirmar(como(client, dueno_a), carga_id)

    tramo = TramoPrecio.objects.de_tienda(tienda_a).get()
    assert tramo.producto_id == producto_a.pk
    assert tramo.valor == 4490


def test_un_tramo_de_un_producto_que_no_existe_se_rechaza(client, dueno_a):
    archivo = xlsx(
        encabezados=["codigo_producto", "umbral", "unidad_umbral", "precio_neto"],
        filas=[["NO-EXISTE", "12", "cantidad", "4490"]],
        nombre="tramos.xlsx",
    )

    cuerpo = validar(como(client, dueno_a), "tramos", archivo).json()

    assert cuerpo["rechazar"] == 1
    assert "catálogo" in cuerpo["rechazos"][0]["motivo"]


# ── Fotos ───────────────────────────────────────────────────────────────────


def test_las_fotos_se_emparejan_por_el_nombre_del_archivo(
    client, dueno_a, tienda_a, producto_a
):
    archivo = zip_archivo({f"{producto_a.codigo}.png": imagen_bytes()})

    cuerpo = validar(como(client, dueno_a), "fotos", archivo).json()
    confirmar(como(client, dueno_a), cuerpo["cargaId"])

    producto_a.refresh_from_db()
    assert cuerpo["emparejadas"] == 1
    assert producto_a.foto


def test_el_nombre_con_que_se_guarda_lo_pone_el_sistema(
    client, dueno_a, producto_a
):
    """El nombre del archivo subido solo sirve para emparejar; no llega al disco."""
    archivo = zip_archivo({f"{producto_a.codigo}.png": imagen_bytes()})
    cuerpo = validar(como(client, dueno_a), "fotos", archivo).json()
    confirmar(como(client, dueno_a), cuerpo["cargaId"])

    producto_a.refresh_from_db()
    assert producto_a.codigo in producto_a.foto.name


def test_el_resumen_dice_qué_quedó_sin_emparejar(client, dueno_a, producto_a):
    archivo = zip_archivo({"no-existe.png": imagen_bytes()})

    cuerpo = validar(como(client, dueno_a), "fotos", archivo).json()

    assert cuerpo["fotosSinProducto"] == ["no-existe.png"]
    assert producto_a.codigo in cuerpo["productosSinFoto"]


def test_un_zip_con_rutas_hacia_afuera_no_escribe_fuera(
    client, dueno_a, producto_a
):
    """
    `../../algo.png` es un intento de escribir donde no corresponde.

    No se rechaza la entrada: se le quita la ruta y queda solo el nombre, que
    es lo único que necesitamos para emparejar.
    """
    archivo = zip_archivo({f"../../{producto_a.codigo}.png": imagen_bytes()})

    cuerpo = validar(como(client, dueno_a), "fotos", archivo).json()
    confirmar(como(client, dueno_a), cuerpo["cargaId"])

    producto_a.refresh_from_db()
    assert cuerpo["emparejadas"] == 1
    assert ".." not in producto_a.foto.name


def test_un_zip_con_demasiados_archivos_se_rechaza(
    client, dueno_a, detalles_de_error, monkeypatch
):
    from apps.ingesta import fotos

    monkeypatch.setattr(fotos, "MAXIMO_ARCHIVOS", 2)
    archivo = zip_archivo({f"f{i}.png": imagen_bytes() for i in range(3)})

    respuesta = validar(como(client, dueno_a), "fotos", archivo)

    assert respuesta.status_code == 400
    assert "máximo" in detalles_de_error(respuesta)["archivo"][0]


def test_un_zip_que_se_descomprime_demasiado_se_rechaza(
    client, dueno_a, detalles_de_error, monkeypatch
):
    """
    Un .zip chico puede descomprimirse en varios gigas y llenar el disco.

    Por eso el tope es sobre el tamaño descomprimido, no sobre el del archivo.
    """
    from apps.ingesta import fotos

    monkeypatch.setattr(fotos, "MAXIMO_DESCOMPRIMIDO", 100)
    archivo = zip_archivo({"grande.png": imagen_bytes(tamano=(200, 200))})

    respuesta = validar(como(client, dueno_a), "fotos", archivo)

    assert respuesta.status_code == 400
    assert "MB" in detalles_de_error(respuesta)["archivo"][0]


def test_un_archivo_que_no_es_imagen_dentro_del_zip_se_rechaza(
    client, dueno_a, producto_a
):
    archivo = zip_archivo({f"{producto_a.codigo}.png": b"esto no es una imagen"})

    cuerpo = validar(como(client, dueno_a), "fotos", archivo).json()

    assert cuerpo["emparejadas"] == 0


def test_un_zip_dañado_se_rechaza(client, dueno_a, detalles_de_error):
    from django.core.files.uploadedfile import SimpleUploadedFile

    roto = SimpleUploadedFile(
        "fotos.zip", b"PK\x03\x04" + b"basura" * 50, content_type="application/zip"
    )
    respuesta = validar(como(client, dueno_a), "fotos", roto)

    assert respuesta.status_code == 400
    assert "archivo" in detalles_de_error(respuesta)


# ── Plantilla, límites e historial ──────────────────────────────────────────


def test_la_plantilla_se_descarga_y_es_un_xlsx(client, dueno_a):
    respuesta = como(client, dueno_a).get("/api/panel/plantillas/productos")

    assert respuesta.status_code == 200
    assert respuesta.content[:2] == b"PK"  # Un .xlsx es un .zip por dentro.


def test_la_plantilla_trae_los_encabezados_que_el_backend_espera(client, dueno_a):
    """
    Si la plantilla y el lector se separan, nadie se entera hasta que falla.

    Esta prueba los ata: lo que se descarga tiene que poder cargarse.
    """
    from openpyxl import load_workbook

    from apps.ingesta import productos as carga_productos

    respuesta = como(client, dueno_a).get("/api/panel/plantillas/productos")
    libro = load_workbook(io.BytesIO(respuesta.content), read_only=True)
    encabezados = [c.value for c in next(libro.active.iter_rows())]

    assert encabezados == carga_productos.COLUMNAS


def test_los_limites_los_dice_el_backend(client, dueno_a, settings):
    settings.TAMANO_MAXIMO_CARGA = 12345

    cuerpo = como(client, dueno_a).get(LIMITES).json()

    assert cuerpo["tamanoMaximoBytes"] == 12345
    assert ".xlsx" in cuerpo["extensiones"]


def test_el_historial_guarda_quién_subió_qué(client, dueno_a, tienda_a):
    validar(como(client, dueno_a), "productos", xlsx())

    linea = como(client, dueno_a).get(HISTORIAL).json()[0]

    assert linea["archivo"] == "catalogo.xlsx"
    assert linea["usuario"] == dueno_a.nombre
    assert linea["tipo"] == "productos"


def test_el_historial_filtra_por_tipo(client, dueno_a, producto_a):
    validar(como(client, dueno_a), "productos", xlsx())
    validar(
        como(client, dueno_a),
        "tramos",
        xlsx(
            encabezados=["codigo_producto", "umbral", "unidad_umbral", "precio_neto"],
            filas=[[producto_a.codigo, "12", "cantidad", "4490"]],
        ),
    )

    cuerpo = como(client, dueno_a).get(HISTORIAL, {"tipo": "tramos"}).json()

    assert [c["tipo"] for c in cuerpo] == ["tramos"]


# ── Quién entra ─────────────────────────────────────────────────────────────


def test_el_vendedor_puede_cargar(client, vendedor_b):
    assert validar(como(client, vendedor_b), "productos", xlsx()).status_code == 201


def test_un_comprador_no_puede_cargar(client, comprador):
    assert validar(como(client, comprador), "productos", xlsx()).status_code == 403


def test_sin_sesion_no_se_carga(client, db):
    assert validar(client, "productos", xlsx()).status_code == 401


# ── Aislamiento ─────────────────────────────────────────────────────────────


@pytest.mark.aislamiento
def test_la_carga_entra_en_la_tienda_del_token(client, dueno_a, tienda_a, tienda_b):
    carga_id = validar(como(client, dueno_a), "productos", xlsx()).json()["cargaId"]
    confirmar(como(client, dueno_a), carga_id)

    assert Producto.objects.de_tienda(tienda_a).count() == 1
    assert not Producto.objects.de_tienda(tienda_b).exists()


@pytest.mark.aislamiento
def test_no_se_confirma_la_carga_de_otra_tienda(
    client, dueno_a, vendedor_b, tienda_b
):
    """
    404 y no 403: el identificador viaja en la URL entre validar y confirmar.
    """
    carga_id = validar(como(client, vendedor_b), "productos", xlsx()).json()["cargaId"]

    respuesta = confirmar(como(client, dueno_a), carga_id)

    assert respuesta.status_code == 404
    assert not Producto.objects.de_tienda(tienda_b).exists()


@pytest.mark.aislamiento
def test_no_se_descarta_la_carga_de_otra_tienda(client, dueno_a, vendedor_b):
    carga_id = validar(como(client, vendedor_b), "productos", xlsx()).json()["cargaId"]

    respuesta = como(client, dueno_a).delete(f"/api/panel/cargas/{carga_id}")

    assert respuesta.status_code == 404
    assert Carga.objects.get(id=carga_id).estado == EstadoCarga.VALIDADA


@pytest.mark.aislamiento
def test_el_historial_de_una_tienda_no_muestra_el_de_la_otra(
    client, dueno_a, vendedor_b
):
    validar(como(client, dueno_a), "productos", xlsx(nombre="de-a.xlsx"))
    validar(como(client, vendedor_b), "productos", xlsx(nombre="de-b.xlsx"))

    de_a = [c["archivo"] for c in como(client, dueno_a).get(HISTORIAL).json()]
    de_b = [c["archivo"] for c in como(client, vendedor_b).get(HISTORIAL).json()]

    assert de_a == ["de-a.xlsx"]
    assert de_b == ["de-b.xlsx"]


@pytest.mark.aislamiento
def test_una_categoria_de_otra_tienda_no_sirve_en_la_planilla(
    client, dueno_a, tienda_b, crear_categoria
):
    """La categoría se busca dentro de la tienda, así que la ajena no existe."""
    crear_categoria(tienda_b, nombre="ajena")
    archivo = xlsx(
        encabezados=["codigo", "nombre", "precio_neto", "categoria"],
        filas=[["CAT-1", "Con categoría ajena", "4990", "ajena"]],
    )

    cuerpo = validar(como(client, dueno_a), "productos", archivo).json()

    assert cuerpo["rechazar"] == 1
    assert "no existe en tu tienda" in cuerpo["rechazos"][0]["motivo"]


@pytest.mark.aislamiento
def test_un_tramo_no_alcanza_el_producto_de_otra_tienda(
    client, dueno_a, producto_b
):
    archivo = xlsx(
        encabezados=["codigo_producto", "umbral", "unidad_umbral", "precio_neto"],
        filas=[[producto_b.codigo, "12", "cantidad", "4490"]],
    )

    cuerpo = validar(como(client, dueno_a), "tramos", archivo).json()

    assert cuerpo["rechazar"] == 1


@pytest.mark.aislamiento
def test_una_foto_no_alcanza_el_producto_de_otra_tienda(
    client, dueno_a, producto_b
):
    archivo = zip_archivo({f"{producto_b.codigo}.png": imagen_bytes()})

    cuerpo = validar(como(client, dueno_a), "fotos", archivo).json()
    confirmar(como(client, dueno_a), cuerpo["cargaId"])

    producto_b.refresh_from_db()
    assert cuerpo["emparejadas"] == 0
    assert not producto_b.foto
