"""
La plantilla que se descarga antes de cargar.

Existe para que nadie tenga que adivinar los nombres de las columnas. Trae los
encabezados exactos, una fila de ejemplo y, donde hay lista cerrada, los
valores que se aceptan escritos en la celda misma.

La fila de ejemplo va marcada: se borra antes de cargar, y si alguien la deja,
la carga la rechaza con su número de fila como cualquier otra.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from apps.catalogo.models import FormaVenta, UnidadMedida
from apps.ingesta import productos, tramos
from apps.ingesta.models import TipoCarga
from apps.precios.models import UnidadUmbral

#: Qué dice cada columna, para la fila de ayuda.
AYUDA = {
    "codigo": "Obligatorio. Como lo llamas tú. No se repite.",
    "nombre": "Obligatorio.",
    "descripcion": "Opcional.",
    "categoria": "Tiene que existir ya en tu tienda.",
    "forma_venta": " | ".join(v for v, _ in FormaVenta.choices),
    "unidad_medida": " | ".join(v for v, _ in UnidadMedida.choices),
    "precio_neto": "Obligatorio. En pesos, sin IVA.",
    "controla_stock": "si | no",
    "stock": "Solo si controla stock.",
    "cantidad_minima": "Opcional. Por omisión, 1.",
    "foto": "Nombre del archivo dentro del .zip de fotos.",
    "codigo_producto": "Obligatorio. El código del producto que ya cargaste.",
    "umbral": "Desde cuánto aplica este precio.",
    "unidad_umbral": " | ".join(v for v, _ in UnidadUmbral.choices),
}

EJEMPLOS = {
    TipoCarga.PRODUCTOS: [
        "EJEMPLO-01",
        "Nombre del producto",
        "Descripción breve",
        "",
        "unidad",
        "unidad",
        "4990",
        "si",
        "10",
        "1",
        "",
    ],
    TipoCarga.TRAMOS: ["EJEMPLO-01", "12", "cantidad", "4490"],
}

COLUMNAS = {
    TipoCarga.PRODUCTOS: productos.COLUMNAS,
    TipoCarga.TRAMOS: tramos.COLUMNAS,
}


def construir(tipo: str) -> bytes:
    """La plantilla de ese tipo de carga, como un .xlsx en memoria."""
    columnas = COLUMNAS[tipo]

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Carga"

    hoja.append(columnas)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill("solid", fgColor="E8EEF4")

    hoja.append([AYUDA.get(c, "") for c in columnas])
    for celda in hoja[2]:
        celda.font = Font(italic=True, color="4A5468", size=9)
        celda.alignment = Alignment(wrap_text=True)

    hoja.append(EJEMPLOS[tipo])
    for celda in hoja[3]:
        celda.font = Font(color="4A5468")

    hoja.append([])
    hoja.append(["← Borra la fila de ayuda y la de ejemplo antes de cargar."])

    for indice, columna in enumerate(columnas, start=1):
        hoja.column_dimensions[hoja.cell(row=1, column=indice).column_letter].width = (
            max(len(columna) + 4, 18)
        )

    memoria = io.BytesIO()
    libro.save(memoria)
    return memoria.getvalue()
